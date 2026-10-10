from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

HEAD = ['#', 'Employee ID', 'Name', 'Department', 'Present', 'Absent', 'Late', 'Leave', '%']
DEPT_HEAD = ['Department', 'Employees', 'Present', 'Absent', 'Late', 'Leave']
NAVY = '0E4A5C'


def _rows(rep):
    return [[i + 1, r['employee_id'], r['name'], r['department'], r['present'],
             r['absent'], r['late'], r['leave'], r['percent']]
            for i, r in enumerate(rep['rows'])]


def _style_header(row):
    for c in row:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=NAVY)


def report_xlsx(rep):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Employee wise'
    ws.append([f"Yahya Hospital - Attendance Report {rep['from']} to {rep['to']}"])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(HEAD)
    _style_header(ws[3])
    for row in _rows(rep):
        ws.append(row)
    for col, w in zip('ABCDEFGHI', [5, 13, 26, 20, 9, 9, 9, 9, 8]):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet('Department wise')
    ws2.append(DEPT_HEAD)
    _style_header(ws2[1])
    for d in rep['departments']:
        ws2.append([d['department'], d['employees'], d['present'], d['absent'], d['late'], d['leave']])
    ws2.column_dimensions['A'].width = 24

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report_pdf(rep):
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=30, rightMargin=30,
                            topMargin=30, bottomMargin=30)
    st = getSampleStyleSheet()
    t = rep['totals']
    els = [
        Paragraph('Yahya Hospital - Attendance Report', st['Title']),
        Paragraph(f"{rep['from']} to {rep['to']} &nbsp;|&nbsp; Present {t['present']} &nbsp; "
                  f"Absent {t['absent']} &nbsp; Late {t['late']} &nbsp; Leave {t['leave']}", st['Normal']),
        Spacer(1, 12),
    ]
    tbl = Table([HEAD] + _rows(rep), repeatRows=1,
                colWidths=[24, 60, 120, 90, 40, 40, 36, 38, 36])
    tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#' + NAVY)),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.lightgrey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F6F8')]),
        ('ALIGN', (4, 0), (-1, -1), 'CENTER'),
    ]))
    els.append(tbl)
    doc.build(els)
    return buf.getvalue()