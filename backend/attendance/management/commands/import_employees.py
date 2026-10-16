import csv
from datetime import datetime
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from attendance.models import AuditLog, Department, Employee, Shift
from attendance.views import make_temp_pin, next_employee_id

HEADERS = ['First Name', 'Last Name', 'Designation', 'Department', 'Phone', 'Shift', 'Role']
ROLES = {'': 'staff', 'staff': 'staff', 'head': 'head', 'department head': 'head', 'hr': 'hr'}


def clean(v):
    s = '' if v is None else str(v).strip()
    return s[:-2] if s.endswith('.0') and s[:-2].isdigit() else s   # Excel turns phone numbers into 3001234567.0


def read_rows(path):
    p = Path(path)
    if not p.exists():
        raise CommandError(f'File not found: {p}')
    ext = p.suffix.lower()
    if ext == '.csv':
        with open(p, newline='', encoding='utf-8-sig') as f:
            return [{(k or '').strip().lower(): clean(v) for k, v in r.items()} for r in csv.DictReader(f)]
    if ext in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        ws = load_workbook(p, read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        head = [clean(h).lower() for h in next(it, [])]
        rows = []
        for r in it:
            if all(clean(c) == '' for c in r):
                continue
            rows.append({head[i]: clean(c) for i, c in enumerate(r) if i < len(head) and head[i]})
        return rows
    raise CommandError('Use a .xlsx or .csv file')


class Command(BaseCommand):
    help = 'Create many employees from an Excel/CSV file. Use --template to get a sample file.'

    def add_arguments(self, parser):
        parser.add_argument('file', nargs='?', help='.xlsx or .csv file to import')
        parser.add_argument('--dry-run', action='store_true', help='check the file, save nothing')
        parser.add_argument('--template', metavar='PATH', help='write a sample .xlsx file and exit')

    def handle(self, *args, **o):
        from openpyxl import Workbook

        if o['template']:
            wb = Workbook()
            ws = wb.active
            ws.append(HEADERS)
            ws.append(['Ali', 'Khan', 'Medical Officer', 'Emergency', '03001234567', '', 'staff'])
            ws.append(['Sara', 'Ahmed', 'Staff Nurse', 'Nursing', '03111234567', '', 'staff'])
            wb.save(o['template'])
            self.stdout.write(self.style.SUCCESS(f'Template written to {o["template"]}'))
            return
        if not o['file']:
            raise CommandError('Give a file path, or use --template sample.xlsx')

        rows = read_rows(o['file'])
        if not rows:
            raise CommandError('No data rows found')
        if 'first name' not in rows[0]:
            raise CommandError('Missing a "First Name" column. Expected columns: ' + ', '.join(HEADERS))

        depts = {d.name.lower(): d for d in Department.objects.all()}
        shifts = {s.name.lower(): s for s in Shift.objects.all()}
        new_depts, errors, plan, seen = {}, [], [], set()

        for i, r in enumerate(rows, start=2):          # row 1 is the header
            first, last = r.get('first name', ''), r.get('last name', '')
            desig, phone = r.get('designation', ''), r.get('phone', '')
            dname, sname = r.get('department', ''), r.get('shift', '')
            role = ROLES.get(r.get('role', '').lower())
            if not first:
                errors.append(f'Row {i}: First Name is empty'); continue
            if role is None:
                errors.append(f'Row {i}: Role "{r.get("role")}" is not valid (use staff, head or hr)'); continue
            if len(phone) > 20:
                errors.append(f'Row {i}: Phone is too long'); continue
            shift = None
            if sname:
                shift = shifts.get(sname.lower())
                if not shift:
                    errors.append(f'Row {i}: Shift "{sname}" does not exist (create it in Settings first)'); continue
            key = (first.lower(), last.lower(), phone or desig.lower())
            if key in seen:
                errors.append(f'Row {i}: {first} {last} appears twice in the file'); continue
            seen.add(key)
            dup = Employee.objects.filter(first_name__iexact=first, last_name__iexact=last)
            dup = dup.filter(phone=phone) if phone else dup.filter(designation__iexact=desig)
            if dup.exists():
                errors.append(f'Row {i}: {first} {last} already exists as {dup.first().username}'); continue
            if dname and dname.lower() not in depts:
                new_depts[dname.lower()] = dname
            plan.append(dict(first=first, last=last, desig=desig, phone=phone, dname=dname, shift=shift, role=role))

        if errors:
            self.stdout.write(self.style.ERROR(f'{len(errors)} problem(s). Nothing was imported:'))
            for e in errors:
                self.stdout.write('  ' + e)
            raise CommandError('Fix the file and run again')

        self.stdout.write(f'{len(plan)} employee(s) are valid.')
        if new_depts:
            self.stdout.write('New departments that will be created (category "other"): ' + ', '.join(new_depts.values()))
        if o['dry_run']:
            self.stdout.write(self.style.SUCCESS('Dry run finished. Nothing was saved.'))
            return

        out = Workbook()
        ws = out.active
        ws.append(['Employee ID', 'Name', 'Department', 'Temporary PIN'])
        with transaction.atomic():
            for key, name in new_depts.items():
                depts[key] = Department.objects.create(name=name, category='other')
            for p in plan:
                pin = make_temp_pin()
                emp = Employee(
                    username=next_employee_id(), first_name=p['first'], last_name=p['last'],
                    designation=p['desig'], phone=p['phone'], shift=p['shift'], role=p['role'],
                    department=depts.get(p['dname'].lower()) if p['dname'] else None,
                    must_change_pin=True,
                )
                emp.set_password(pin)
                emp.save()
                ws.append([emp.username, f'{p["first"]} {p["last"]}'.strip(), p['dname'], pin])
            AuditLog.objects.create(actor=None, action='employee.bulk_import', target=Path(o['file']).name,
                                    details={'count': len(plan)})

        result = f'import_result_{datetime.now():%Y%m%d_%H%M%S}.xlsx'
        out.save(result)
        self.stdout.write(self.style.SUCCESS(f'Imported {len(plan)} employees.'))
        self.stdout.write(f'Employee IDs and temporary PINs saved in {result}. Hand them out, then delete that file.')